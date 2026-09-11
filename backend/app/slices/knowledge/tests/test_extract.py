import io
from pathlib import Path

from app.slices.knowledge.extract import extract_text


def _build_minimal_pdf_bytes() -> bytes:
    """A single-page PDF built with pypdf's own writer.

    A hand-written PDF relying on pypdf reconstructing a missing xref table
    (strict=False) turned out to raise `PdfReadError: startxref not found`
    on the installed pypdf version, since that fixture has no `startxref`
    keyword at all for the recovery path to find. Building the fixture with
    PdfWriter instead sidesteps that entirely: no fixture library, and the
    file pypdf produces is one pypdf itself can always parse back.
    """
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    writer = PdfWriter()
    page = writer.add_blank_page(width=300, height=144)

    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 18 Tf 20 60 Td (refund policy thirty days) Tj ET")
    stream_ref = writer._add_object(stream)

    font = DictionaryObject()
    font[NameObject("/Type")] = NameObject("/Font")
    font[NameObject("/Subtype")] = NameObject("/Type1")
    font[NameObject("/BaseFont")] = NameObject("/Helvetica")
    font_ref = writer._add_object(font)

    font_dict = DictionaryObject()
    font_dict[NameObject("/F1")] = font_ref
    resources = DictionaryObject()
    resources[NameObject("/Font")] = font_dict

    page[NameObject("/Resources")] = resources
    page[NameObject("/Contents")] = stream_ref

    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


MINIMAL_PDF = _build_minimal_pdf_bytes()


def test_plain_text_passthrough(tmp_path: Path):
    source = tmp_path / "notes.md"
    source.write_text("# Hours\nOpen 9 to 5", encoding="utf-8")
    assert extract_text(source, "text/markdown") == "# Hours\nOpen 9 to 5"


def test_pdf_text_is_extracted(tmp_path: Path):
    source = tmp_path / "policy.pdf"
    source.write_bytes(MINIMAL_PDF)
    text = extract_text(source, "application/pdf")
    assert "refund policy thirty days" in text


def test_docx_paragraphs_are_extracted(tmp_path: Path):
    import docx

    document = docx.Document()
    document.add_paragraph("Shipping takes three days.")
    document.add_paragraph("Returns are free.")
    source = tmp_path / "shipping.docx"
    document.save(str(source))

    text = extract_text(
        source,
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    assert "Shipping takes three days." in text
    assert "Returns are free." in text


def test_html_tags_are_stripped(tmp_path: Path):
    source = tmp_path / "faq.html"
    source.write_text("<h1>FAQ</h1><p>We ship <b>worldwide</b>.</p>", encoding="utf-8")
    text = extract_text(source, "text/html")
    assert "<" not in text
    assert "FAQ" in text and "We ship" in text and "worldwide" in text


def test_extension_wins_when_content_type_is_generic(tmp_path: Path):
    source = tmp_path / "policy.pdf"
    source.write_bytes(MINIMAL_PDF)
    assert "refund policy" in extract_text(source, "application/octet-stream")
